from openpyxl import load_workbook
from pathlib import Path
from copy import copy
import zipfile


class WorkbookWriter:
    """在一次同步中复用已打开的工作簿，最后统一保存。"""

    def __init__(self, file):
        self.file = Path(file)

        if not self.file.exists():
            raise FileNotFoundError(f"文件不存在:{self.file}")

        if not zipfile.is_zipfile(self.file):
            raise Exception(f"Excel文件损坏:{self.file}")

        self.wb = load_workbook(self.file)
        self.changed = False

    def write_sheet(self, sheet, data, start_row=4):
        if sheet not in self.wb.sheetnames:
            return False

        ws = self.wb[sheet]



        # generator转列表，用于获取插入行数
        rows = list(data)



        if not rows:
            return True



        column_count = max(ws.max_column, len(rows[0]))
        template_height = ws.row_dimensions[start_row].height

        # 在表头下面插入新行
        ws.insert_rows(start_row, amount=len(rows))

        # insert_rows 不会继承单元格样式。插入后位于新行下方的原数据行
        # 仍保留模板格式，因此用它作为样式样本填充所有新行。
        template_row = start_row + len(rows)
        for row_number in range(start_row, template_row):
            self._copy_row_style(
                ws, template_row, row_number, column_count, template_height
            )

        current_row = start_row
        for item in rows:
            for col, value in enumerate(item.values(), start=1):
                ws.cell(row=current_row, column=col, value=value)
            current_row += 1

        self.changed = True
        return True

    @staticmethod
    def _copy_row_style(
            ws,
            source_row,
            target_row,
            column_count,
            template_height
    ):
        for column in range(1, column_count + 1):
            ws.cell(target_row, column)._style = copy(ws.cell(source_row, column)._style)

        source_dimension = ws.row_dimensions[source_row]
        target_dimension = ws.row_dimensions[target_row]
        target_dimension.height = template_height or source_dimension.height
        target_dimension.hidden = source_dimension.hidden

    def find_header_row(self, sheet, headers, max_row=20):
        """查找从第 1 列开始、与 headers 完全一致的表头行。"""
        if sheet not in self.wb.sheetnames:
            return None

        ws = self.wb[sheet]
        for row_number in range(1, min(ws.max_row, max_row) + 1):
            values = [ws.cell(row_number, column).value for column in range(1, len(headers) + 1)]
            if values == headers:
                return row_number

        return None

    def save(self):
        if self.changed:
            self.wb.save(self.file)

    def close(self):
        self.wb.close()


def write_sheet(file, sheet, data, start_row=4):
    """兼容旧调用：打开、写入、保存、关闭。"""
    writer = WorkbookWriter(file)
    try:
        result = writer.write_sheet(sheet, data, start_row)
        writer.save()
        return result
    finally:
        writer.close()
